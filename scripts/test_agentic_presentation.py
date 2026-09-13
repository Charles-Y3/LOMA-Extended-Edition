# -*- coding: utf-8 -*-
from pipeline.deliverables.presentation_finalize import (
    finalize_presentation_source,
    is_compile_ready_presentation,
    presentation_chat_summary,
)
from pipeline.deliverables.presentation_prepare import prepare_agentic_presentation
from pipeline.deliverables.presentation_compile import compile_agentic_presentation
from pptx import Presentation


class _Plan:
    steps = []


PROSE = """
Presentation Structure (5 Slide Themes):

What is Kindness? – Define kindness through definitions and benefits.
The Ripple Effect – Highlight data showing how kindness spreads.
Types of Kindness – Categorize kindness into self, others, environment.
How to Practice Kindness Daily – Offer simple actionable steps.
Why Kindness Matters – Conclude with well-being message.
"""

md = finalize_presentation_source({}, _Plan(), PROSE, "create a 5 slides beautiful presentation on kindness")
print("compile_ready", is_compile_ready_presentation(md))
print("--- markdown ---")
print(md)
prep = prepare_agentic_presentation(md, workflow_instruction="create a 5 slides beautiful presentation on kindness")
print("prep_ok", prep.ok, prep.errors)
path = compile_agentic_presentation(prep.markdown, "output", prep.theme, query="kindness")
prs = Presentation(path)
print("pptx_slides", len(prs.slides))
print("slide0_title", prs.slides[0].shapes.title.text)
print("--- chat ---")
print(presentation_chat_summary(prep.markdown, "output_loma.pptx"))
