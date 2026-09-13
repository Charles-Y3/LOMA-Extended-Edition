# tools/file_parser.py
# -*- coding: utf-8 -*-
import os
import base64


def _collapse_vertical_cjk_lines(text):
    """Rejoin one-character-per-line PDF text extraction into flowing text.

    For vertically-set / columnar CJK PDFs, PyMuPDF's and pypdf's text extraction put
    each glyph on its own line (one row per character in the page's visual layout) —
    the character SEQUENCE is still correct reading order, but a translator model sees
    what looks like a wall of single-character "sentences" and struggles with it
    (verified: stripping all whitespace from the raw extraction reproduces the source
    docx's paragraph text exactly, in order — this is purely a line-break artifact, not
    a reordering bug). Whitespace-only lines mark real paragraph/sentence gaps in the
    source layout, so those become paragraph breaks; everything else is concatenated
    with no separator. Horizontal PDFs (normal multi-character lines) are detected and
    left untouched.
    """
    if not text:
        return text
    lines = text.split("\n")
    non_empty = [ln for ln in lines if ln.strip()]
    if len(non_empty) < 8:
        return text
    short = sum(1 for ln in non_empty if len(ln.strip()) <= 2)
    if short / len(non_empty) < 0.6:
        return text
    out = []
    for ln in lines:
        if not ln.strip():
            if out and out[-1] != "\n\n":
                out.append("\n\n")
            continue
        out.append(ln)
    return "".join(out)


def parse_uploaded_file(filepath):
    """
    Reads an uploaded file and returns a standardized dictionary.

    Prefer services.source_parser.parse_file / attach_uploaded_file for new code.
    This module remains the low-level byte → markdown extractor used by source_parser.
    """
    ext = os.path.splitext(filepath)[1].lower()
    filename = os.path.basename(filepath)

    # 1. Plain Text & Code Formats
    if ext in ['.txt', '.md', '.csv', '.json', '.py', '.html']:
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                return {
                    "filename": filename,
                    "type": "text",
                    "content": f.read()
                }
        except Exception as e:
            return {"filename": filename, "type": "error", "content": f"Text read error: {str(e)}"}

    # 2. Portable Document Format (PDF) Extraction
    elif ext == '.pdf':
        try:
            extracted_text = ""
            try:
                # PyMuPDF decodes embedded CJK/CID fonts far more reliably than pypdf,
                # which falls back to raw (often pinyin-named) glyph names when a PDF's
                # ToUnicode CMap is missing/broken, corrupting Chinese text on extraction.
                import pymupdf as fitz
                with fitz.open(filepath) as doc:
                    for page_num, page in enumerate(doc):
                        page_text = _collapse_vertical_cjk_lines(page.get_text()).strip()
                        if page_text:
                            extracted_text += f"--- Page {page_num + 1} ---\n{page_text}\n"
            except ImportError:
                import pypdf
                reader = pypdf.PdfReader(filepath)
                for page_num, page in enumerate(reader.pages):
                    page_text = _collapse_vertical_cjk_lines(page.extract_text()).strip()
                    if page_text:
                        extracted_text += f"--- Page {page_num + 1} ---\n{page_text}\n"

            if not extracted_text.strip():
                return {"filename": filename, "type": "error",
                        "content": "PDF extracted as empty text (could be scanned image/requires OCR)."}

            return {
                "filename": filename,
                "type": "text",
                "content": extracted_text
            }
        except Exception as e:
            return {"filename": filename, "type": "error", "content": f"PDF parse error: {str(e)}"}

    # 3. Microsoft Word (DOCX) Extraction - Structured Format Preservation
    elif ext == '.docx':
        try:
            import docx

            doc = docx.Document(filepath)
            extracted_text = ""

            def parse_run_formatting(run):
                """Wraps raw text runs with basic markdown typography tags."""
                text = run.text
                if not text.strip():
                    return text
                if run.bold and run.italic:
                    return f"***{text}***"
                elif run.bold:
                    return f"**{text}**"
                elif run.italic:
                    return f"*{text}*"
                return text

            def parse_paragraph_to_markdown(p):
                """Maps document styles and lists directly into standard markdown equivalents."""
                p_text = "".join(parse_run_formatting(r) for r in p.runs).strip()
                if not p_text:
                    return ""

                style_name = p.style.name.lower()

                # Check Header Styles
                if style_name.startswith('heading 1'):
                    return f"# {p_text}\n\n"
                elif style_name.startswith('heading 2'):
                    return f"## {p_text}\n\n"
                elif style_name.startswith('heading 3'):
                    return f"### {p_text}\n\n"
                elif style_name.startswith('heading 4'):
                    return f"#### {p_text}\n\n"

                # Check Bullet/List Styles
                elif 'list bullet' in style_name or p.style.name.startswith('List Bullet'):
                    return f"* {p_text}\n"
                elif 'list number' in style_name or p.style.name.startswith('List Number'):
                    return f"1. {p_text}\n"

                # Default Plain Paragraph text block
                return f"{p_text}\n\n"

            def parse_table_to_markdown(table):
                """Converts matrix block data arrays into structured Markdown tables."""
                table_md = ""
                for row_idx, row in enumerate(table.rows):
                    row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_md += "| " + " | ".join(row_text) + " |\n"
                    if row_idx == 0:
                        # Append the required alignment/demarcation boundary rule
                        table_md += "| " + " | ".join(["---"] * len(row_text)) + " |\n"
                return table_md + "\n"

            # Parse elements sequentially in their true document order layout
            for element in doc.element.body:
                if element.tag.endswith('p'):
                    p = docx.text.paragraph.Paragraph(element, doc)
                    extracted_text += parse_paragraph_to_markdown(p)
                elif element.tag.endswith('tbl'):
                    t = docx.table.Table(element, doc)
                    extracted_text += parse_table_to_markdown(t)

            if not extracted_text.strip():
                return {"filename": filename, "type": "text", "content": ""}

            return {"filename": filename, "type": "text", "content": extracted_text.strip()}

        except ImportError:
            return {
                "filename": filename,
                "type": "error",
                "content": "The 'python-docx' library is missing. Run 'pip install python-docx' to parse Word files locally."
            }
        except Exception as e:
            return {"filename": filename, "type": "error", "content": f"DOCX parse error: {str(e)}"}

    # 4. Microsoft PowerPoint (PPTX) Extraction - Slide & Table Layout Preservation
    elif ext == '.pptx':
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            prs = Presentation(filepath)
            slides_markdown = []

            for slide_idx, slide in enumerate(prs.slides, start=1):
                # Explicitly mapped slide boundary for the output generator to anchor onto
                slide_elements = [f"--- Slide {slide_idx} ---"]

                if slide.shapes.title and slide.shapes.title.text.strip():
                    slide_elements.append(f"## {slide.shapes.title.text.strip()}")

                for shape in slide.shapes:
                    if shape == slide.shapes.title:
                        continue

                    if hasattr(shape, "text_frame") and shape.text_frame:
                        text_blocks = []
                        for paragraph in shape.text_frame.paragraphs:
                            p_text = paragraph.text.strip()
                            if p_text:
                                if paragraph.level > 0:
                                    indent = "  " * paragraph.level
                                    text_blocks.append(f"{indent}* {p_text}")
                                else:
                                    text_blocks.append(p_text)
                        if text_blocks:
                            slide_elements.append("\n".join(text_blocks))

                    elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                        table = shape.table
                        table_md = ""
                        for row_idx in range(len(table.rows)):
                            cells_text = [table.cell(row_idx, col_idx).text.strip().replace("\n", " ")
                                          for col_idx in range(len(table.columns))]
                            table_md += "| " + " | ".join(cells_text) + " |\n"
                            if row_idx == 0:
                                table_md += "| " + " | ".join(["---"] * len(cells_text)) + " |\n"
                        slide_elements.append(table_md)

                slides_markdown.append("\n\n".join(slide_elements))

            if not slides_markdown:
                return {"filename": filename, "type": "error",
                        "content": "PowerPoint document parsed as empty presentation layout."}

            return {
                "filename": filename,
                "type": "text",
                "content": "\n\n".join(slides_markdown)
            }
        except ImportError:
            return {
                "filename": filename,
                "type": "error",
                "content": "The 'python-pptx' library is missing. Run 'pip install python-pptx' to parse PowerPoint files locally."
            }
        except Exception as e:
            return {"filename": filename, "type": "error", "content": f"PPTX parse error: {str(e)}"}

    # 5. Spreadsheet
    elif ext in [".xlsx", ".xls"]:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            parts = [f"# Spreadsheet: {filename}"]

            def _row_cells(row) -> list[str]:
                return ["" if c is None else str(c).strip() for c in row]

            def _looks_like_header(cells: list[str]) -> bool:
                filled = [c for c in cells if c]
                if len(filled) < 2:
                    return False
                # Skip Excel filter/dropdown placeholder rows (single label, no data columns).
                if len(filled) == 1:
                    return False
                alpha = sum(1 for c in filled if any(ch.isalpha() for ch in c))
                return alpha >= max(1, len(filled) // 3)

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                parts.append(f"\n## Sheet: {sheet_name}\n")
                raw_rows = [_row_cells(row) for row in ws.iter_rows(values_only=True)]
                raw_rows = [r for r in raw_rows if any(c for c in r)]
                header_idx = 0
                for i, cells in enumerate(raw_rows[:15]):
                    if _looks_like_header(cells):
                        header_idx = i
                        break
                data_rows = raw_rows[header_idx:]
                rows = ["| " + " | ".join(cells) + " |" for cells in data_rows]
                if rows:
                    ncol = max(len(r.split("|")) - 2 for r in rows)
                    parts.append("| " + " | ".join(["---"] * max(1, ncol)) + " |")
                    parts.extend(rows)
            wb.close()
            body = "\n".join(parts).strip()
            return {"filename": filename, "type": "text", "content": body or "(empty spreadsheet)"}
        except ImportError:
            return {
                "filename": filename,
                "type": "error",
                "content": "Install openpyxl to parse Excel files: pip install openpyxl",
            }
        except Exception as e:
            return {"filename": filename, "type": "error", "content": f"XLSX parse error: {str(e)}"}

    # 6. Audio — defer transcription; chat uses omni model or on-demand Whisper install
    elif ext in [".mp3", ".wav", ".m4a", ".ogg", ".flac"]:
        return {
            "filename": filename,
            "type": "media_audio",
            "content": os.path.abspath(filepath),
        }

    # 7. Video — defer transcription; vision/audio models or Whisper later
    elif ext in [".mp4", ".mkv", ".mov", ".webm", ".avi"]:
        return {
            "filename": filename,
            "type": "media_video",
            "content": os.path.abspath(filepath),
        }

    # 8. Multimodal Image Formats
    elif ext in ['.png', '.jpg', '.jpeg', '.webp']:
        return {
            "filename": filename,
            "type": "image",
            "content": os.path.abspath(filepath)
        }

    else:
        return {
            "filename": filename,
            "type": "unsupported",
            "content": f"Unsupported file extension: {ext}"
        }