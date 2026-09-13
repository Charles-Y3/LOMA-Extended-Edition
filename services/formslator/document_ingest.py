# -*- coding: utf-8 -*-
"""Convert uploaded documents to Formslator-ready .docx."""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from typing import Any, Callable

from services.formslator.paths import UPLOADS_DIR, ensure_dirs


class PasswordRequired(Exception):
    """Raised when a password is needed to open the document."""

    def __init__(self, path: str, *, kind: str = "document") -> None:
        self.path = path
        self.kind = kind
        super().__init__(f"Password required for {kind}")


def _out_docx_path(src: str, suffix: str = "_prepared") -> str:
    base = os.path.splitext(os.path.basename(src))[0]
    return os.path.join(UPLOADS_DIR, f"{base}{suffix}.docx")


def _pdf_needs_password(path: str) -> bool:
    import pypdf

    try:
        reader = pypdf.PdfReader(path)
        return bool(getattr(reader, "is_encrypted", False))
    except Exception:
        return False


def _docx_needs_password(path: str) -> bool:
    try:
        import msoffcrypto

        with open(path, "rb") as f:
            office = msoffcrypto.OfficeFile(f)
            return bool(office.is_encrypted())
    except ImportError:
        try:
            from docx import Document

            Document(path)
            return False
        except Exception as exc:
            msg = str(exc).lower()
            if "password" in msg or "encrypt" in msg:
                return True
            raise
    except Exception:
        return False


def inspect_upload(path: str) -> str | None:
    """Return 'pdf' or 'docx' if password required, else None."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf" and _pdf_needs_password(path):
        return "pdf"
    if ext in (".docx", ".doc") and _docx_needs_password(path):
        return "docx"
    return None


def _pdf_decrypt_temp(path: str, password: str) -> str:
    """Write an unencrypted temp PDF for Office converters; caller must delete."""
    import pypdf

    reader = pypdf.PdfReader(path)
    if getattr(reader, "is_encrypted", False):
        if reader.decrypt(password) == 0:
            raise ValueError("Incorrect PDF password.")
    writer = pypdf.PdfWriter()
    writer.append(reader)
    fd, tmp = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)
    try:
        with open(tmp, "wb") as f:
            writer.write(f)
    except Exception:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise
    return tmp


def _pdf_via_word(abs_pdf: str, out: str, *, timeout_s: float = 25.0) -> bool:
    """Try Word PDF→DOCX in a worker thread; abandon if it hangs (common on PDF open)."""
    import sys
    import threading

    if sys.platform != "win32":
        return False

    holder: dict[str, Any] = {"ok": False, "word": None}

    def _run() -> None:
        try:
            import pythoncom
            import win32com.client

            pythoncom.CoInitialize()
            word = None
            try:
                word = win32com.client.DispatchEx("Word.Application")
                holder["word"] = word
                word.Visible = False
                word.DisplayAlerts = 0
                try:
                    word.Options.ConfirmConversions = False
                except Exception:
                    pass
                doc = word.Documents.Open(
                    FileName=abs_pdf,
                    ConfirmConversions=False,
                    ReadOnly=True,
                    AddToRecentFiles=False,
                    Visible=False,
                )
                doc.SaveAs2(out, FileFormat=16)
                doc.Close(False)
                holder["ok"] = os.path.isfile(out)
            finally:
                if word is not None:
                    try:
                        word.Quit()
                    except Exception:
                        pass
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass
        except Exception:
            holder["ok"] = False

    t = threading.Thread(target=_run, daemon=True)
    t.start()
    t.join(timeout_s)
    if t.is_alive():
        # Word is still blocked (often on PDF open). Best-effort quit; then fall through.
        try:
            w = holder.get("word")
            if w is not None:
                w.Quit()
        except Exception:
            pass
        return False
    return bool(holder.get("ok"))


def _pdf_via_libreoffice(abs_pdf: str, out: str, *, timeout_s: float = 45.0) -> bool:
    out_dir = tempfile.mkdtemp()
    try:
        for cmd in (
            ["soffice", "--headless", "--nologo", "--nofirststartwizard",
             "--convert-to", "docx", "--outdir", out_dir, abs_pdf],
            ["libreoffice", "--headless", "--nologo", "--nofirststartwizard",
             "--convert-to", "docx", "--outdir", out_dir, abs_pdf],
        ):
            exe = shutil.which(cmd[0])
            if not exe:
                continue
            try:
                subprocess.run(
                    [exe, *cmd[1:]],
                    check=True,
                    capture_output=True,
                    timeout=timeout_s,
                )
            except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
                continue
            candidate = os.path.join(
                out_dir, os.path.splitext(os.path.basename(abs_pdf))[0] + ".docx"
            )
            if os.path.isfile(candidate):
                os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
                if candidate != out:
                    shutil.move(candidate, out)
                return True
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)
    return False


_PDF_FONT_ALIASES: dict[str, str] = {
    "DFYanKai": "華康正顏楷體",
    "DFFangSong": "華康仿宋體",
    "DFWeiBei": "華康魏碑體",
    "DFKaiShu": "華康楷書",
    "DFKaiSho": "華康楷書",
    "DFHeiMedium": "華康中黑體",
    "DFHeiBold": "華康粗黑體",
    "DFMing": "華康明體",
    "DFLiShu": "華康隸書",
    "MicrosoftJhengHei": "微軟正黑體",
    "MicrosoftYaHei": "微软雅黑",
    "PMingLiU": "新細明體",
    "MingLiU": "細明體",
    "KaiTi": "楷体",
    "FangSong": "仿宋",
    "SimSun": "宋体",
    "SimHei": "黑体",
    "NotoSansCJK": "Noto Sans CJK",
    "NotoSerifCJK": "Noto Serif CJK",
}


def _clean_pdf_font_name(raw: str | None) -> str:
    """Normalize PDF BaseFont to a friendly display/signature name."""
    import re

    if not raw:
        return ""
    name = str(raw).strip().lstrip("/")
    if "+" in name:
        name = name.split("+", 1)[1]
    for suffix in (
        "-WIN-BF",
        "-HK-BF",
        "-BF",
        "-Identity-H",
        "-Identity-V",
        "-Estd",
        "-SB",
    ):
        if name.endswith(suffix):
            name = name[: -len(suffix)]
    name = re.sub(r"-W\d+$", "", name)
    name = re.sub(r"-(Md|Bold|Light|Regular|Medium|SemiBold|SB)$", "", name, flags=re.I)
    name = re.sub(r"Regular$", "", name, flags=re.I)
    name = name.strip(" -_") or "?"

    # Prefer known Chinese aliases (subset fonts from Word→PDF).
    for key, alias in _PDF_FONT_ALIASES.items():
        if name == key or name.startswith(key):
            return alias
    return name


def _round_pdf_font_size(font_size: float | None) -> float:
    """Snap PDF float sizes (14.04, 15.96) to mapping-friendly points."""
    try:
        v = float(font_size or 0)
    except Exception:
        return 0.0
    if v <= 0:
        return 0.0
    return float(round(v))


def _set_run_font(run, font_name: str, font_size: float | None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    if font_name and font_name != "?":
        run.font.name = font_name
        try:
            rpr = run._element.get_or_add_rPr()
            r_fonts = rpr.get_or_add_rFonts()
            r_fonts.set(qn("w:ascii"), font_name)
            r_fonts.set(qn("w:hAnsi"), font_name)
            r_fonts.set(qn("w:eastAsia"), font_name)
        except Exception:
            pass
    if font_size and font_size > 0:
        try:
            run.font.size = Pt(round(float(font_size) * 2) / 2)
        except Exception:
            pass


def _pdf_font_spans_to_docx(path: str, *, password: str | None = None) -> str | None:
    """Rebuild a .docx from PDF text spans with font/size so Format mapping can detect styles."""
    import pypdf
    from docx import Document

    reader = pypdf.PdfReader(path)
    if getattr(reader, "is_encrypted", False):
        if not password:
            raise PasswordRequired(path, kind="PDF")
        if reader.decrypt(password) == 0:
            raise ValueError("Incorrect PDF password.")

    # (page, -y, x, text, font, size)
    spans: list[tuple[int, float, float, str, str, float]] = []
    for page_i, page in enumerate(reader.pages):
        page_spans: list[tuple[float, float, str, str, float]] = []

        def visitor(text, cm, tm, font_dict, font_size, _ps=page_spans) -> None:
            t = text or ""
            if not t.strip():
                return
            try:
                x = float(tm[4])
                y = float(tm[5])
            except Exception:
                x, y = 0.0, 0.0
            fname = None
            if font_dict:
                fname = font_dict.get("/BaseFont") or font_dict.get("/FontName")
                try:
                    fname = str(fname)
                except Exception:
                    fname = None
            try:
                size = _round_pdf_font_size(font_size)
            except Exception:
                size = 0.0
            _ps.append((y, x, t, _clean_pdf_font_name(fname), size))

        try:
            page.extract_text(visitor_text=visitor)
        except TypeError:
            try:
                page.extract_text(visitor_oper=visitor)
            except Exception:
                continue
        except Exception:
            continue

        page_spans.sort(key=lambda s: (-round(s[0], 1), s[1]))
        for y, x, t, font, size in page_spans:
            spans.append((page_i, -y, x, t, font, size))

    if not spans:
        return None

    # Group into paragraphs by vertical gaps; merge adjacent same-font runs.
    paras: list[list[tuple[str, str, float]]] = []
    cur: list[tuple[str, str, float]] = []
    prev_key: tuple[int, float] | None = None
    y_break = 3.0  # PDF points — new para when line jumps

    for page_i, neg_y, _x, text, font, size in spans:
        y = -neg_y
        key = (page_i, round(y, 1))
        if prev_key is None:
            cur = [(text, font, size)]
            prev_key = (page_i, y)
        else:
            same_page = page_i == prev_key[0]
            dy = abs(y - prev_key[1]) if same_page else 999.0
            if (not same_page) or dy > y_break:
                if cur:
                    paras.append(cur)
                cur = [(text, font, size)]
            else:
                last_t, last_f, last_s = cur[-1]
                if last_f == font and abs(last_s - size) < 0.6:
                    cur[-1] = (last_t + text, last_f, last_s)
                else:
                    cur.append((text, font, size))
            prev_key = (page_i, y)
    if cur:
        paras.append(cur)

    # Drop empty / digit-only paragraphs (same filter as detect_styles)
    cleaned: list[list[tuple[str, str, float]]] = []
    for runs in paras:
        text = "".join(t for t, _, _ in runs).strip()
        if not text or text.strip().rstrip(".").isdigit():
            continue
        cleaned.append(runs)
    if not cleaned:
        return None

    out = _out_docx_path(path, "_from_pdf_fonts")
    doc = Document()
    for runs in cleaned:
        para = doc.add_paragraph()
        for text, font, size in runs:
            if not text:
                continue
            run = para.add_run(text)
            _set_run_font(run, font, size if size > 0 else None)
    doc.save(out)

    # Must yield at least one detectable signature for Format mapping.
    try:
        from services.formslator.style_service import detect_styles

        if not detect_styles(out):
            return None
    except Exception:
        return None
    return out


def _pdf_text_extract_to_docx(path: str, *, password: str | None = None) -> str:
    import pypdf
    from docx import Document

    reader = pypdf.PdfReader(path)
    if getattr(reader, "is_encrypted", False):
        if not password:
            raise PasswordRequired(path, kind="PDF")
        if reader.decrypt(password) == 0:
            raise ValueError("Incorrect PDF password.")
    chunks: list[str] = []
    for page in reader.pages:
        chunks.append((page.extract_text() or "").strip())
    text = "\n\n".join(c for c in chunks if c)
    if not text.strip():
        raise ValueError("Could not extract text from PDF (may be scanned images).")
    out = _out_docx_path(path, "_from_pdf")
    doc = Document()
    for block in text.split("\n\n"):
        block = block.strip()
        if block:
            doc.add_paragraph(block)
    doc.save(out)
    return out


def _pdf_to_docx(path: str, *, password: str | None = None) -> str:
    """Convert PDF → DOCX: LibreOffice, then Word (timed), then plain text extract.

    LibreOffice is tried first because Word COM frequently hangs indefinitely on PDF open.
    """
    ensure_dirs()
    abs_pdf = os.path.abspath(path)
    tmp_pdf: str | None = None

    if _pdf_needs_password(path):
        if not password:
            raise PasswordRequired(path, kind="PDF")
        tmp_pdf = _pdf_decrypt_temp(path, password)
        abs_pdf = tmp_pdf

    office_out = os.path.abspath(_out_docx_path(path, "_from_pdf_office"))
    try:
        # LibreOffice first — has a real process timeout; Word PDF open often never returns.
        if _pdf_via_libreoffice(abs_pdf, office_out):
            try:
                from services.formslator.style_service import detect_styles

                if detect_styles(office_out):
                    return office_out
            except Exception:
                pass
        if _pdf_via_word(abs_pdf, office_out):
            try:
                from services.formslator.style_service import detect_styles

                if detect_styles(office_out):
                    return office_out
            except Exception:
                pass
    finally:
        if tmp_pdf:
            try:
                os.remove(tmp_pdf)
            except OSError:
                pass

    # Font/size-aware rebuild (Word-exported PDFs) so Format mapping has style buckets.
    font_doc = _pdf_font_spans_to_docx(path, password=password)
    if font_doc:
        return font_doc

    return _pdf_text_extract_to_docx(path, password=password)


def _decrypt_docx(path: str, password: str) -> str:
    try:
        import io

        import msoffcrypto
        from docx import Document

        decrypted = io.BytesIO()
        with open(path, "rb") as f:
            office = msoffcrypto.OfficeFile(f)
            office.load_key(password=password)
            office.decrypt(decrypted)
        decrypted.seek(0)
        out = _out_docx_path(path, "_decrypted")
        Document(decrypted).save(out)
        return out
    except ImportError as exc:
        raise RuntimeError(
            "Encrypted Office files require: pip install msoffcrypto-tool"
        ) from exc
    except Exception as exc:
        msg = str(exc).lower()
        if "password" in msg or "decrypt" in msg:
            raise ValueError("Incorrect document password.") from exc
        raise


def _doc_to_docx(path: str) -> str:
    import shutil
    import subprocess
    import sys
    import tempfile

    # Word's SaveAs2 (and soffice --outdir matching below) reject a relative
    # output path — must be absolute, same as office_out in _pdf_to_docx.
    out = os.path.abspath(_out_docx_path(path, "_from_doc"))
    abs_doc = os.path.abspath(path)

    errors: list[str] = []

    if sys.platform == "win32":
        try:
            import pythoncom
            import win32com.client

            pythoncom.CoInitialize()
            word = None
            try:
                word = win32com.client.DispatchEx("Word.Application")
                word.Visible = False
                word.DisplayAlerts = 0
                doc = word.Documents.Open(FileName=abs_doc, ReadOnly=True, Visible=False)
                doc.SaveAs2(out, FileFormat=16)
                doc.Close(False)
                return out
            finally:
                if word is not None:
                    word.Quit()
                pythoncom.CoUninitialize()
        except ImportError:
            errors.append("Word automation needs pywin32 (Microsoft Word installed).")
        except Exception as exc:
            errors.append(f"Word conversion failed: {exc}")

    out_dir = tempfile.mkdtemp()
    try:
        found_soffice = False
        for cmd in (
            ["soffice", "--headless", "--convert-to", "docx", "--outdir", out_dir, abs_doc],
            ["libreoffice", "--headless", "--convert-to", "docx", "--outdir", out_dir, abs_doc],
        ):
            exe = shutil.which(cmd[0])
            if not exe:
                continue
            found_soffice = True
            try:
                subprocess.run(cmd[:1] + cmd[1:], check=True, capture_output=True, timeout=120)
            except Exception as exc:
                errors.append(f"LibreOffice conversion failed: {exc}")
                continue
            candidate = os.path.join(
                out_dir, os.path.splitext(os.path.basename(abs_doc))[0] + ".docx"
            )
            if os.path.isfile(candidate):
                if candidate != out:
                    shutil.move(candidate, out)
                return out
        if not found_soffice:
            errors.append("LibreOffice not found on PATH.")
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)

    detail = " ".join(errors) or "no converter available."
    raise RuntimeError(
        f"Could not convert .doc — install Microsoft Word or LibreOffice, or save as .docx first. ({detail})"
    )


def prepare_input_document(path: str, *, password: str | None = None) -> str:
    """Return a .docx path usable by Formslator."""
    ensure_dirs()
    if not os.path.isfile(path):
        raise FileNotFoundError(path)
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        if inspect_upload(path) == "docx":
            if not password:
                raise PasswordRequired(path, kind="Word document")
            return _decrypt_docx(path, password)
        return path
    if ext == ".pdf":
        return _pdf_to_docx(path, password=password)
    if ext == ".doc":
        return _doc_to_docx(path)
    raise ValueError(f"Unsupported document type: {ext}")


def prompt_password_dialog(
    *,
    kind: str,
    on_submit: Callable[[str], None],
) -> None:
    from nicegui import ui
    from pipeline.i18n import t as tr

    kind_label = {
        "pdf": tr("formslator.upload.kind_pdf"),
        "docx": tr("formslator.upload.kind_docx"),
        "doc": tr("formslator.upload.kind_docx"),
        "word document": tr("formslator.upload.kind_docx"),
        "PDF": tr("formslator.upload.kind_pdf"),
    }.get((kind or "").strip().lower(), tr("formslator.upload.kind_document"))

    with ui.dialog() as dialog, ui.card().classes("p-4 gap-3 min-w-[280px]"):
        ui.label(tr("formslator.upload.password_title", kind=kind_label)).classes(
            "text-sm font-semibold"
        )
        pwd = ui.input(tr("formslator.upload.password_label"), password=True).props(
            "dense dark standout"
        ).classes("w-full")

        def submit() -> None:
            value = (pwd.value or "").strip()
            if not value:
                ui.notify(tr("formslator.upload.password_required"), type="warning")
                return
            dialog.close()
            on_submit(value)

        with ui.row().classes("w-full justify-end gap-2"):
            ui.button(tr("common.cancel"), on_click=dialog.close).props("flat dense")
            ui.button(tr("formslator.upload.unlock"), on_click=submit).props(
                "flat dense color=primary"
            )
    dialog.open()
