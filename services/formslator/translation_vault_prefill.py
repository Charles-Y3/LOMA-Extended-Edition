# -*- coding: utf-8 -*-
"""Prefill Formslator translation column from Document Intelligence vault."""
from __future__ import annotations

import re
from typing import Callable

from docx import Document
from docx.shared import RGBColor

from extensions.document_intelligence.translation import translation_index_for_scope
from services.formslator.format_engine import column_indices
from services.formslator.translate_engine import (
    clean_basic_grammar,
    extract_nonempty_paragraphs,
    to_title_case,
)
from services.formslator.vault_alignment import (
    ALIGN_STRICT,
    collect_vault_candidates,
    is_valid_vault_output,
    search_vault_translation,
)
from pipeline.i18n import t as tr

LogFn = Callable[[str], None]
VAULT_GREEN = RGBColor(0x18, 0x6A, 0x3B)
_LINE_SPLIT = re.compile(r"\n+")


def prefill_translations_from_vault(
    formatted_doc_path: str,
    *,
    original_column: str = "left",
    output_doc_path: str | None = None,
    log_fn: LogFn | None = None,
    vault_scope: str = "all",
    vault_library_ids: list[str] | None = None,
) -> tuple[int, int]:
    """Write vault matches into the translation column (green text for matches)."""
    log = log_fn or (lambda _m: None)
    index = translation_index_for_scope(vault_scope, vault_library_ids)
    pair_count = len(index.pairs)
    if pair_count == 0:
        log(tr("formslator.log.vault_empty"))
        return 0, 0

    log(tr("formslator.log.vault_pairs", count=pair_count))

    if output_doc_path is None:
        output_doc_path = formatted_doc_path

    orig_idx, trans_idx = column_indices(original_column)
    doc = Document(formatted_doc_path)
    paragraphs = extract_nonempty_paragraphs(formatted_doc_path, original_column)

    rows_all = []
    for table in doc.tables:
        for row in table.rows:
            rows_all.append(row)

    matched = 0
    total = sum(1 for p in paragraphs if (p or "").strip())

    for idx, row in enumerate(rows_all):
        orig_cell = row.cells[orig_idx]
        trans_cell = row.cells[trans_idx]
        source = paragraphs[idx] if idx < len(paragraphs) else ""
        if not (source or "").strip():
            continue

        if not collect_vault_candidates(index, source):
            continue

        translation, is_vault, method = _match_paragraph(index, source)
        if not translation or not is_vault:
            _clear_translation_cell(trans_cell)
            log(
                f"Vault skipped (no valid English extraction): "
                f"{source[:60]}{'…' if len(source) > 60 else ''}"
            )
            continue

        if trans_cell.paragraphs:
            trans_para = trans_cell.paragraphs[0]
        else:
            trans_para = trans_cell.add_paragraph("")

        _write_translation_cell(
            trans_para,
            translation,
            vault_match=True,
            orig_cell=orig_cell,
        )
        matched += 1
        log(
            tr(
                "formslator.log.vault_match",
                method=method,
                snippet=f"{source[:60]}{'…' if len(source) > 60 else ''}",
            )
        )

    doc.save(output_doc_path)
    log(tr("formslator.log.vault_prefill", matched=matched, total=total))
    return matched, total


def _match_paragraph(index, text: str) -> tuple[str, bool, str]:
    """Return (translation, is_vault_match, method)."""
    stripped = (text or "").strip()
    if not stripped:
        return "", False, ""

    translation, ok, method = _vault_translation_for_phrase(index, stripped)
    if ok:
        return translation, True, method

    lines = [ln.strip() for ln in _LINE_SPLIT.split(stripped) if ln.strip()]
    if len(lines) <= 1:
        return "", False, ""

    out_lines: list[str] = []
    methods: list[str] = []
    matched_any = False
    for line in lines:
        line_tr, line_ok, line_method = _vault_translation_for_phrase(index, line)
        if line_ok:
            out_lines.append(line_tr)
            methods.append(line_method)
            matched_any = True
        else:
            out_lines.append("")
    if not matched_any:
        return "", False, ""
    method = methods[0] if len(set(methods)) == 1 else "mixed"
    return "\n".join(out_lines), True, method


def _vault_translation_for_phrase(index, phrase: str) -> tuple[str, bool, str]:
    hit = search_vault_translation(
        phrase,
        index,
        min_alignment=ALIGN_STRICT,
        allow_llm=True,
        borderline_llm=True,
        llm_validate_borderline=True,
    )
    if not hit or not is_valid_vault_output(phrase, hit.translation_text):
        return "", False, ""
    return hit.translation_text, True, hit.method


def _clear_translation_cell(trans_cell) -> None:
    if trans_cell.paragraphs:
        trans_para = trans_cell.paragraphs[0]
    else:
        trans_para = trans_cell.add_paragraph("")
    for r in list(trans_para.runs):
        r._element.getparent().remove(r._element)
    trans_para.add_run("")


def _write_translation_cell(
    trans_para,
    translation_text: str,
    *,
    vault_match: bool,
    orig_cell,
) -> None:
    for r in list(trans_para.runs):
        r._element.getparent().remove(r._element)

    if not translation_text:
        trans_para.add_run("")
        return

    style_name = trans_para.style.name.strip() if trans_para.style else ""
    txt_to_process = translation_text
    if style_name and style_name.startswith("E1"):
        txt_to_process = to_title_case(txt_to_process)
    elif len(txt_to_process) > 1:
        txt_to_process = txt_to_process[0].upper() + txt_to_process[1:]
    elif txt_to_process:
        txt_to_process = txt_to_process.upper()
    txt_to_process = clean_basic_grammar(txt_to_process)

    run = trans_para.add_run(txt_to_process)
    if vault_match:
        run.font.color.rgb = VAULT_GREEN

    if orig_cell.paragraphs and orig_cell.paragraphs[0].alignment is not None:
        trans_para.alignment = orig_cell.paragraphs[0].alignment
