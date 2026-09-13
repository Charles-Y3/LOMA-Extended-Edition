# -*- coding: utf-8 -*-
"""Source language detection for Formslator."""
from __future__ import annotations

import re
from typing import Any

from docx import Document

CJK_RE = re.compile(r"[\u3400-\u9FFF\uF900-\uFAFF]")
THAI_RE = re.compile(r"[\u0E00-\u0E7F]")
ARABIC_RE = re.compile(r"[\u0600-\u06FF]")
HIRAGANA_KATAKANA_RE = re.compile(r"[\u3040-\u30FF]")

CUSTOM_CODE = "__custom__"

OUTPUT_LANGUAGES: list[tuple[str, str]] = [
    ("English", "en"),
    ("Spanish", "es"),
    ("French", "fr"),
    ("German", "de"),
    ("Portuguese", "pt"),
    ("Italian", "it"),
    ("Vietnamese", "vi"),
    ("Thai", "th"),
    ("Indonesian", "id"),
    ("Malay", "ms"),
    ("Tagalog", "tl"),
    ("Hindi", "hi"),
    ("Arabic", "ar"),
    ("Korean", "ko"),
    ("Japanese", "ja"),
    ("Simplified Chinese", "zh_cn"),
    ("Traditional Chinese", "zh_tw"),
]

LANG_NAMES: dict[str, str] = {
    "en": "English",
    "es": "Spanish",
    "fr": "French",
    "de": "German",
    "pt": "Portuguese",
    "it": "Italian",
    "vi": "Vietnamese",
    "th": "Thai",
    "id": "Indonesian",
    "ms": "Malay",
    "tl": "Tagalog",
    "hi": "Hindi",
    "ar": "Arabic",
    "ko": "Korean",
    "ja": "Japanese",
    "zh_cn": "Simplified Chinese",
    "zh_tw": "Traditional Chinese",
    "zh": "Chinese",
    "unknown": "the source language",
    CUSTOM_CODE: "Custom",
}

CJK_TARGETS = {"zh_cn", "zh_tw", "ja", "ko"}

# Model-family heuristics → preferred target codes (hints only; full list still available).
_MODEL_LANG_HINTS: list[tuple[tuple[str, ...], list[str]]] = [
    (("qwen", "yi-", "deepseek", "internlm", "chatglm", "glm-", "hunyuan"),
     ["zh_cn", "zh_tw", "en", "ja", "ko"]),
    (("llama", "mistral", "mixtral", "phi", "vicuna", "zephyr"),
     ["en", "es", "fr", "de", "pt", "it"]),
    (("gemma",),
     ["en", "es", "fr", "de", "ja", "ko"]),
    (("aya", "command-r", "command_r"),
     ["en", "es", "fr", "de", "pt", "ar", "ja", "zh_cn", "ko", "hi"]),
]


def recommended_language_codes(model: str = "") -> list[str]:
    """Best-effort preferred target codes for a model tag (not a hard filter)."""
    name = (model or "").lower()
    known = {code for _, code in OUTPUT_LANGUAGES}
    for needles, codes in _MODEL_LANG_HINTS:
        if any(n in name for n in needles):
            return [c for c in codes if c in known]
    return [c for c in ("en", "es", "zh_cn", "zh_tw", "ja") if c in known]


def output_language_select_options(model: str = "") -> dict[str, str]:
    """Localized labels for Formslator output-language dropdown (recommended first)."""
    from pipeline.i18n import t as tr

    by_code = {code: tr(f"formslator.lang.{code}") for _, code in OUTPUT_LANGUAGES}
    rec = recommended_language_codes(model)
    rest = [c for _, c in OUTPUT_LANGUAGES if c not in rec]
    ordered: dict[str, str] = {}
    for code in rec:
        ordered[by_code[code]] = code
    for code in rest:
        ordered[by_code[code]] = code
    ordered[tr("formslator.lang.custom")] = CUSTOM_CODE
    return ordered


def resolve_target_name(target_code: str, custom_name: str = "") -> str:
    """Display / prompt name for a target language code."""
    if target_code == CUSTOM_CODE:
        return (custom_name or "").strip() or "the target language"
    return LANG_NAMES.get(target_code, target_code)


def sample_text_from_docx(path: str, max_chars: int = 4000) -> str:
    try:
        doc = Document(path)
    except Exception:
        return ""

    chunks: list[str] = []
    total = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = p.text.strip()
                    if t:
                        chunks.append(t)
                        total += len(t)
                        if total >= max_chars:
                            return "\n".join(chunks)[:max_chars]
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            chunks.append(t)
            total += len(t)
            if total >= max_chars:
                break
    return "\n".join(chunks)[:max_chars]


def detect_source_language(text: str) -> str:
    if not text.strip():
        return "unknown"

    n = max(len(text), 1)
    cjk = len(CJK_RE.findall(text))
    thai = len(THAI_RE.findall(text))
    arabic = len(ARABIC_RE.findall(text))
    kana = len(HIRAGANA_KATAKANA_RE.findall(text))

    if cjk / n > 0.15 and kana / max(cjk, 1) > 0.05:
        return "ja"
    if cjk / n > 0.15:
        return "zh"
    if thai / n > 0.15:
        return "th"
    if arabic / n > 0.15:
        return "ar"
    return "en"


def detect_source_from_docx(path: str) -> str:
    return detect_source_language(sample_text_from_docx(path))


def target_is_cjk(code: str) -> bool:
    return code in CJK_TARGETS


def script_info(lang_code: str) -> tuple[re.Pattern[str], str]:
    """Return regex and display label for character counting in final statistics."""
    code = (lang_code or "unknown").lower()
    if code in ("zh", "zh_cn", "zh_tw"):
        return CJK_RE, "Chinese"
    if code == "ja":
        return re.compile(r"[\u3040-\u30FF\u3400-\u9FFF\uF900-\uFAFF]"), "Japanese"
    if code == "ko":
        return re.compile(r"[\uAC00-\uD7AF\u1100-\u11FF]"), "Korean"
    if code == "th":
        return THAI_RE, "Thai"
    if code == "ar":
        return ARABIC_RE, "Arabic"
    if code == "hi":
        return re.compile(r"[\u0900-\u097F]"), "Hindi"
    if code in ("en", "es", "fr", "de", "pt", "it", "vi", "id", "ms", "tl", "unknown", CUSTOM_CODE):
        return re.compile(r"\S"), "text"
    name = LANG_NAMES.get(code, code)
    return re.compile(r"\S"), name


def count_lang_chars(text: str, lang_code: str) -> int:
    pattern, _ = script_info(lang_code)
    return len(pattern.findall(text or ""))


def count_doc_lang_chars(path: str, lang_code: str) -> int:
    """Count language-appropriate characters across all paragraphs and table cells."""
    try:
        doc = Document(path)
    except Exception:
        return 0
    pattern, _ = script_info(lang_code)
    total = 0
    for p in doc.paragraphs:
        total += len(pattern.findall(p.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                total += len(pattern.findall(cell.text))
    return total


def stats_char_label(lang_code: str) -> str:
    _, label = script_info(lang_code)
    if label == "text":
        return "characters"
    return f"{label} chars"


def language_meta(
    source_code: str,
    target_code: str,
    *,
    target_name: str | None = None,
) -> dict[str, Any]:
    return {
        "source_code": source_code,
        "target_code": target_code,
        "source_name": LANG_NAMES.get(source_code, source_code),
        "target_name": (target_name or "").strip()
        or resolve_target_name(target_code),
        "is_cjk_target": target_is_cjk(target_code),
    }
