# -*- coding: utf-8 -*-
"""Apply mutation text maps back into Office files."""
from __future__ import annotations

import os
import re
from typing import Callable

from services.office_mutation.unit_enrich import changed_text_map
from services.office_mutation.unit_style import apply_run_font_snapshot, snapshot_run_font
from services.session import state

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from docx import Document
except ImportError:
    Document = None


def _strip_office_markup(text: str) -> str:
    out = (text or "").strip()
    out = re.sub(r"\*\*(.+?)\*\*", r"\1", out)
    out = re.sub(r"__(.+?)__", r"\1", out)
    out = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", out)
    out = re.sub(r"(?<!\w)_(.+?)_(?!\w)", r"\1", out)
    return out


def write_paragraph_text(paragraph, new_text: str) -> None:
    """Replace paragraph text; keep first run styling (avoids run-split corruption)."""
    new_text = _strip_office_markup(new_text or "")
    if not new_text:
        return
    runs = list(paragraph.runs)
    if runs:
        snap = snapshot_run_font(runs[0])
        runs[0].text = new_text
        for run in runs[1:]:
            run.text = ""
        apply_run_font_snapshot(runs[0], snap)
        return
    paragraph.text = new_text


def _apply_pptx_shape(shapes, slide_idx: int, path_prefix: str, id_to_text: dict[str, str], stats: dict) -> None:
    for shape_idx, shape in enumerate(shapes):
        shape_key = f"{path_prefix}sh{shape_idx}"
        if shape.has_text_frame:
            for para_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                if not paragraph.text.strip():
                    continue
                uid = f"s{slide_idx}_{shape_key}_p{para_idx}"
                if uid in id_to_text:
                    write_paragraph_text(paragraph, id_to_text[uid])
                    stats["applied"] += 1
        if shape.has_table:
            for row_idx, row in enumerate(shape.table.rows):
                for col_idx, cell in enumerate(row.cells):
                    for para_idx, paragraph in enumerate(cell.text_frame.paragraphs):
                        if not paragraph.text.strip():
                            continue
                        uid = f"s{slide_idx}_{shape_key}_t{row_idx}_{col_idx}_p{para_idx}"
                        if uid in id_to_text:
                            write_paragraph_text(paragraph, id_to_text[uid])
                            stats["applied"] += 1
        if getattr(shape, "shape_type", None) == 6:
            try:
                _apply_pptx_shape(shape.shapes, slide_idx, f"{shape_key}_grp_", id_to_text, stats)
            except Exception:
                pass


def apply_pptx(dest_path: str, units: list[dict], text_map: dict[str, str], on_slide_applied: Callable | None = None) -> str:
    if not Presentation:
        state.add_log("python-pptx not installed.")
        return ""

    prs = Presentation(dest_path)
    patch = changed_text_map(units, text_map)
    if not patch:
        state.add_log("Mutation apply: no text changes to patch.")
        return os.path.abspath(dest_path)
    id_to_text = {uid: _strip_office_markup(val) for uid, val in patch.items()}
    stats = {"applied": 0}
    total_slides = len(prs.slides)

    for slide_idx, slide in enumerate(prs.slides):
        _apply_pptx_shape(slide.shapes, slide_idx, "", id_to_text, stats)
        prs.save(dest_path)
        if on_slide_applied:
            try:
                on_slide_applied(slide_idx, total_slides, dest_path)
            except Exception:
                pass

    return os.path.abspath(dest_path)


def apply_docx(dest_path: str, units: list[dict], text_map: dict[str, str]) -> str:
    if not Document:
        state.add_log("python-docx not installed.")
        return ""

    doc = Document(dest_path)
    patch = changed_text_map(units, text_map)
    if not patch:
        state.add_log("Mutation apply: no text changes to patch.")
        return os.path.abspath(dest_path)
    id_to_text = {uid: _strip_office_markup(val) for uid, val in patch.items()}

    for para_idx, paragraph in enumerate(doc.paragraphs):
        uid = f"p{para_idx}"
        if uid in id_to_text and paragraph.text.strip():
            write_paragraph_text(paragraph, id_to_text[uid])

    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                for para_idx, paragraph in enumerate(cell.text_frame.paragraphs):
                    uid = f"t{table_idx}_r{row_idx}_c{col_idx}_p{para_idx}"
                    if uid in id_to_text and paragraph.text.strip():
                        write_paragraph_text(paragraph, id_to_text[uid])
        if table_idx % 3 == 0:
            doc.save(dest_path)

    doc.save(dest_path)
    force_docx_portrait(dest_path)
    return os.path.abspath(dest_path)


def force_docx_portrait(path: str) -> None:
    if not Document:
        return
    try:
        from docx.enum.section import WD_ORIENT

        doc = Document(path)
        for section in doc.sections:
            if section.orientation == WD_ORIENT.LANDSCAPE:
                new_width, new_height = section.page_height, section.page_width
                section.orientation = WD_ORIENT.PORTRAIT
                section.page_width = new_width
                section.page_height = new_height
            else:
                section.orientation = WD_ORIENT.PORTRAIT
        doc.save(path)
    except Exception:
        pass


def apply_xlsx(
    dest_path: str,
    units: list[dict],
    text_map: dict[str, str],
    on_unit_applied: Callable | None = None,
) -> str:
    """Apply cell value mutations to an xlsx workbook."""
    try:
        import openpyxl
    except ImportError:
        state.add_log("openpyxl required for spreadsheet mutation.")
        return ""
    if not os.path.isfile(dest_path):
        return ""
    wb = openpyxl.load_workbook(dest_path)
    patch = changed_text_map(units, text_map)
    if not patch:
        state.add_log("Mutation apply: no cell changes to patch.")
        return os.path.abspath(dest_path)
    id_to_val = {uid: _strip_office_markup(val) for uid, val in patch.items()}
    changed = 0
    for u in units:
        uid = u["id"]
        if uid not in id_to_val:
            continue
        new_val = id_to_val[uid]
        ref = u.get("ref") or ()
        if len(ref) < 4 or ref[0] != "cell":
            continue
        sheet_idx, row_idx, col_idx = ref[1], ref[2], ref[3]
        if sheet_idx >= len(wb.worksheets):
            continue
        sheet = wb.worksheets[sheet_idx]
        cell = sheet.cell(row=row_idx + 1, column=col_idx + 1)
        if new_val.startswith("="):
            cell.value = new_val
        else:
            try:
                if "." in new_val and new_val.replace(".", "", 1).isdigit():
                    cell.value = float(new_val)
                elif new_val.isdigit():
                    cell.value = int(new_val)
                else:
                    cell.value = new_val
            except Exception:
                cell.value = new_val
        changed += 1
        if on_unit_applied:
            on_unit_applied(changed, len(units), dest_path)
    wb.save(dest_path)
    state.add_log(f"Spreadsheet mutation applied ({changed} cells).")
    return os.path.abspath(dest_path)
