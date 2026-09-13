# -*- coding: utf-8 -*-
"""Extract mutable text units with content, structure, and style metadata."""
from __future__ import annotations

import os

from services.office_mutation.unit_enrich import (
    detect_lang_hint,
    enrich_neighbor_context,
    heading_level_from_style,
)
from services.office_mutation.unit_style import paragraph_run_styles

try:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
except ImportError:
    Presentation = None
    PP_ALIGN = None

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None
    WD_ALIGN_PARAGRAPH = None


def _alignment_name(value) -> str:
    if value is None:
        return ""
    try:
        return str(value).split(".")[-1].lower()
    except Exception:
        return str(value)


def _pptx_para_alignment(paragraph) -> str:
    try:
        if paragraph.alignment is not None:
            return _alignment_name(paragraph.alignment)
    except Exception:
        pass
    return ""


def _docx_para_alignment(paragraph) -> str:
    try:
        if paragraph.paragraph_format.alignment is not None:
            return _alignment_name(paragraph.paragraph_format.alignment)
    except Exception:
        pass
    return ""


def _docx_list_level(paragraph) -> int | None:
    try:
        lvl = paragraph.paragraph_format.left_indent
        if lvl is not None and lvl:
            return int(lvl.pt) if hasattr(lvl, "pt") else None
    except Exception:
        pass
    return None


def _docx_section_index(doc, para_idx: int) -> int:
    """Approximate section index from paragraph position."""
    try:
        sections = doc.sections
        if len(sections) <= 1:
            return 0
        total = len(doc.paragraphs)
        if total <= 0:
            return 0
        ratio = para_idx / max(total, 1)
        return min(int(ratio * len(sections)), len(sections) - 1)
    except Exception:
        return 0


def _finish_unit(unit: dict) -> dict:
    text = unit.get("text") or unit.get("value") or ""
    unit["lang_hint"] = detect_lang_hint(text)
    if "value" not in unit:
        unit["value"] = text
    return unit


def _collect_pptx_paragraphs(
    shapes,
    slide_idx: int,
    units: list,
    path_prefix: str,
    *,
    title_shape=None,
    slide_number: int = 1,
) -> None:
    for shape_idx, shape in enumerate(shapes):
        shape_key = f"{path_prefix}sh{shape_idx}"
        shape_name = ""
        shape_type = ""
        try:
            shape_name = (shape.name or "").strip()
            shape_type = str(getattr(shape, "shape_type", "") or "")
        except Exception:
            pass

        if shape.has_text_frame:
            for para_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                text = paragraph.text.strip()
                if not text:
                    continue
                uid = f"s{slide_idx}_{shape_key}_p{para_idx}"
                if title_shape is not None and shape == title_shape and para_idx == 0:
                    role = "slide_title"
                elif para_idx == 0 and shape_idx == 0 and not any(
                    u.get("role") == "slide_title"
                    for u in units
                    if (u.get("slide") or 0) == slide_number
                ):
                    role = "slide_title"
                elif para_idx == 0:
                    role = "shape_heading"
                else:
                    role = "body"
                units.append(
                    _finish_unit(
                        {
                            "id": uid,
                            "text": text,
                            "format": "presentation",
                            "slide": slide_number,
                            "role": role,
                            "ref": ("paragraph", slide_idx, shape_key, para_idx),
                            "shape": {
                                "index": shape_idx,
                                "name": shape_name,
                                "type": shape_type,
                            },
                            "location": {
                                "slide": slide_number,
                                "shape_index": shape_idx,
                                "paragraph": para_idx,
                            },
                            "style": {
                                "alignment": _pptx_para_alignment(paragraph),
                                "runs": paragraph_run_styles(paragraph),
                            },
                        }
                    )
                )
        if shape.has_table:
            for row_idx, row in enumerate(shape.table.rows):
                for col_idx, cell in enumerate(row.cells):
                    for para_idx, paragraph in enumerate(cell.text_frame.paragraphs):
                        text = paragraph.text.strip()
                        if not text:
                            continue
                        uid = f"s{slide_idx}_{shape_key}_t{row_idx}_{col_idx}_p{para_idx}"
                        units.append(
                            _finish_unit(
                                {
                                    "id": uid,
                                    "text": text,
                                    "format": "presentation",
                                    "slide": slide_number,
                                    "role": "table_cell",
                                    "ref": (
                                        "table_cell",
                                        slide_idx,
                                        shape_key,
                                        row_idx,
                                        col_idx,
                                        para_idx,
                                    ),
                                    "shape": {
                                        "index": shape_idx,
                                        "name": shape_name,
                                        "type": shape_type,
                                    },
                                    "structure": {"row": row_idx, "col": col_idx},
                                    "location": {
                                        "slide": slide_number,
                                        "shape_index": shape_idx,
                                        "row": row_idx,
                                        "col": col_idx,
                                        "paragraph": para_idx,
                                    },
                                    "style": {
                                        "alignment": _pptx_para_alignment(paragraph),
                                        "runs": paragraph_run_styles(paragraph),
                                    },
                                }
                            )
                        )
        if getattr(shape, "shape_type", None) == 6:
            try:
                _collect_pptx_paragraphs(
                    shape.shapes,
                    slide_idx,
                    units,
                    f"{shape_key}_grp_",
                    title_shape=title_shape,
                    slide_number=slide_number,
                )
            except Exception:
                pass


def extract_pptx_units(source_path: str) -> list[dict]:
    if not Presentation or not source_path.lower().endswith(".pptx"):
        return []
    prs = Presentation(source_path)
    units: list[dict] = []
    for slide_idx, slide in enumerate(prs.slides):
        title_shape = getattr(slide.shapes, "title", None)
        _collect_pptx_paragraphs(
            slide.shapes,
            slide_idx,
            units,
            "",
            title_shape=title_shape,
            slide_number=slide_idx + 1,
        )
    enrich_neighbor_context(units, group_key=lambda u: u.get("slide"))
    return units


def _docx_para_unit(
    doc,
    *,
    para_idx: int,
    paragraph,
    table_idx: int | None = None,
    row_idx: int | None = None,
    col_idx: int | None = None,
) -> dict:
    text = paragraph.text.strip()
    style_name = ""
    try:
        if paragraph.style and paragraph.style.name:
            style_name = paragraph.style.name
    except Exception:
        pass
    heading_level = heading_level_from_style(style_name)
    section_idx = _docx_section_index(doc, para_idx)

    if table_idx is not None:
        uid = f"t{table_idx}_r{row_idx}_c{col_idx}_p{para_idx}"
        ref = ("doc_table", table_idx, row_idx, col_idx, para_idx)
        location = {
            "block": "table",
            "table": table_idx,
            "row": row_idx,
            "col": col_idx,
            "paragraph": para_idx,
        }
    else:
        uid = f"p{para_idx}"
        ref = ("doc_para", para_idx)
        location = {"block": "body", "paragraph": para_idx}

    return _finish_unit(
        {
            "id": uid,
            "text": text,
            "format": "document",
            "page": section_idx + 1,
            "section": section_idx,
            "style_name": style_name,
            "heading_level": heading_level,
            "alignment": _docx_para_alignment(paragraph),
            "list_level": _docx_list_level(paragraph),
            "ref": ref,
            "location": location,
            "style": {"runs": paragraph_run_styles(paragraph)},
        }
    )


def extract_docx_units(source_path: str) -> list[dict]:
    if not Document or not source_path.lower().endswith(".docx"):
        return []
    doc = Document(source_path)
    units: list[dict] = []
    for para_idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            units.append(_docx_para_unit(doc, para_idx=para_idx, paragraph=paragraph))
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                for para_idx, paragraph in enumerate(cell.paragraphs):
                    text = paragraph.text.strip()
                    if text:
                        units.append(
                            _docx_para_unit(
                                doc,
                                para_idx=para_idx,
                                paragraph=paragraph,
                                table_idx=table_idx,
                                row_idx=row_idx,
                                col_idx=col_idx,
                            )
                        )
    enrich_neighbor_context(units, group_key=lambda u: (u.get("section"), u.get("location", {}).get("block")))
    return units


def _xlsx_cell_style(cell) -> dict:
    style: dict = {}
    try:
        if cell.number_format:
            style["number_format"] = cell.number_format
    except Exception:
        pass
    try:
        font = cell.font
        style["font"] = {
            "name": font.name or "",
            "bold": bool(font.bold),
            "italic": bool(font.italic),
            "size": font.size,
            "color": str(font.color.rgb) if font.color and font.color.rgb else "",
        }
    except Exception:
        pass
    try:
        align = cell.alignment
        style["alignment"] = {
            "horizontal": str(align.horizontal or ""),
            "vertical": str(align.vertical or ""),
            "wrap_text": bool(align.wrap_text),
        }
    except Exception:
        pass
    return style


def extract_xlsx_units(source_path: str) -> list[dict]:
    if not source_path.lower().endswith((".xlsx", ".xlsm")):
        return []
    try:
        import openpyxl
    except ImportError:
        return []
    if not os.path.isfile(source_path):
        return []
    wb = openpyxl.load_workbook(source_path, data_only=False)
    units: list[dict] = []
    for sheet_idx, sheet in enumerate(wb.worksheets):
        sheet_name = sheet.title or f"Sheet{sheet_idx + 1}"
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and cell.data_type != "f":
                    continue
                val = cell.value
                if val is None:
                    continue
                text = str(val).strip()
                if not text:
                    continue
                col_letter = cell.column_letter
                address = f"{col_letter}{cell.row}"
                uid = f"sh{sheet_idx}_{col_letter}{cell.row}"
                kind = "formula" if isinstance(val, str) and val.startswith("=") else "data"
                if cell.row == 1:
                    kind = "header"
                units.append(
                    _finish_unit(
                        {
                            "id": uid,
                            "text": text,
                            "value": text,
                            "format": "spreadsheet",
                            "kind": kind,
                            "sheet": {"index": sheet_idx, "name": sheet_name},
                            "cell": {
                                "address": address,
                                "row": cell.row,
                                "col": cell.column,
                                "col_letter": col_letter,
                            },
                            "ref": ("cell", sheet_idx, cell.row - 1, cell.column - 1),
                            "style": _xlsx_cell_style(cell),
                        }
                    )
                )
    enrich_neighbor_context(
        units,
        group_key=lambda u: (u.get("sheet") or {}).get("index"),
    )
    return units


def extract_units(source_path: str) -> list[dict]:
    """Extract all mutable text units from an Office or PDF file."""
    if (source_path or "").lower().endswith(".pdf"):
        from services.office_mutation.pdf_units import extract_pdf_units

        return extract_pdf_units(source_path)
    if not source_path or not os.path.isfile(source_path):
        return []
    ext = os.path.splitext(source_path)[1].lower()
    if ext == ".pptx":
        return extract_pptx_units(source_path)
    if ext == ".docx":
        return extract_docx_units(source_path)
    if ext in (".xlsx", ".xlsm"):
        return extract_xlsx_units(source_path)
    return []
