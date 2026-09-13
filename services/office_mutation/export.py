# -*- coding: utf-8 -*-
"""Write mutated Office output (including cross-format export)."""
from __future__ import annotations

import os
from collections import defaultdict

from services.office_mutation.apply import force_docx_portrait


def _slide_index(unit: dict) -> int:
    ref = unit.get("ref") or ()
    if len(ref) > 1 and isinstance(ref[1], int):
        return ref[1]
    uid = unit.get("id") or ""
    m = __import__("re").match(r"^s(\d+)_", uid)
    return int(m.group(1)) if m else 0


def export_pptx_units_to_docx(units: list[dict], text_map: dict[str, str], dest_path: str) -> str:
    """
    Build Word doc from pptx mutation units preserving slide structure:
    - Heading 1 per slide (title or first block)
    - Body paragraphs / bullets under each slide section
    """
    try:
        from docx import Document
    except ImportError:
        return ""

    by_slide: dict[int, list[dict]] = defaultdict(list)
    for unit in units:
        by_slide[_slide_index(unit)].append(unit)

    doc = Document()
    for slide_idx in sorted(by_slide.keys()):
        slide_units = by_slide[slide_idx]
        title_unit = next((u for u in slide_units if u.get("role") == "slide_title"), None)
        if not title_unit and slide_units:
            title_unit = slide_units[0]

        if title_unit:
            title = (text_map.get(title_unit["id"]) or title_unit.get("text") or "").strip()
            if title:
                doc.add_heading(title, level=1)

        for unit in slide_units:
            if title_unit and unit["id"] == title_unit["id"]:
                continue
            text = (text_map.get(unit["id"]) or unit.get("text") or "").strip()
            if not text:
                continue
            if unit.get("role") == "shape_heading":
                doc.add_heading(text, level=2)
            else:
                for line in text.splitlines():
                    line = line.strip()
                    if line:
                        doc.add_paragraph(line)

    os.makedirs(os.path.dirname(dest_path) or ".", exist_ok=True)
    doc.save(dest_path)
    force_docx_portrait(dest_path)
    return os.path.abspath(dest_path)


def export_docx_from_units(units: list[dict], text_map: dict[str, str], dest_path: str) -> str:
    """Flat export fallback (non-pptx sources)."""
    try:
        from docx import Document
    except ImportError:
        return ""

    if units and any((u.get("id") or "").startswith("s") for u in units):
        return export_pptx_units_to_docx(units, text_map, dest_path)

    doc = Document()
    seen: set[str] = set()
    for unit in units:
        uid = unit["id"]
        if uid in seen:
            continue
        seen.add(uid)
        text = (text_map.get(uid) or unit.get("text") or "").strip()
        if text:
            doc.add_paragraph(text)
    os.makedirs(os.path.dirname(dest_path) or ".", exist_ok=True)
    doc.save(dest_path)
    force_docx_portrait(dest_path)
    return os.path.abspath(dest_path)
