# -*- coding: utf-8 -*-
"""Convert office files and markdown drafts into styled semantic HTML for markup view."""
from __future__ import annotations

import html
import os
import re

from services.presentation_markdown import normalize_presentation_markdown, split_presentation_slides

_EXTENSION_BY_OUTPUT: dict[str, str] = {
    "document": ".docx",
    "presentation": ".pptx",
    "spreadsheet": ".xlsx",
    "pdf": ".pdf",
}


def _escape(text: str) -> str:
    return html.escape(text or "")


def _sanitize_office_html(html: str) -> str:
    """Strip inline widths that break responsive markup panes (mammoth/Word output)."""
    if not html:
        return html
    cleaned = re.sub(r'\s+width="[^"]*"', "", html, flags=re.IGNORECASE)

    def _clean_style(match: re.Match) -> str:
        style = match.group(1)
        parts = [
            p.strip()
            for p in style.split(";")
            if p.strip() and not p.strip().lower().startswith("width")
        ]
        if not parts:
            return ""
        return f' style="{"; ".join(parts)}"'

    return re.sub(r'\sstyle="([^"]*)"', _clean_style, cleaned, flags=re.IGNORECASE)


def _apply_inline_markdown(text: str) -> str:
    """Bold/italic for fallback renderer when python-markdown is unavailable."""
    out = _escape(text or "")
    out = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", out)
    out = re.sub(r"__(.+?)__", r"<strong>\1</strong>", out)
    out = re.sub(r"\*(.+?)\*", r"<em>\1</em>", out)
    out = re.sub(r"(?<!\w)_(.+?)_(?!\w)", r"<em>\1</em>", out)
    return out


def _basic_markdown_to_html_body(text: str) -> str:
    """Lightweight markdown → HTML (headings, lists, emphasis, raw HTML blocks)."""
    raw = text or ""
    if not raw.strip():
        return "<p></p>"

    parts: list[str] = []
    in_ul = False
    para_lines: list[str] = []

    def flush_para() -> None:
        nonlocal para_lines
        if not para_lines:
            return
        joined = " ".join(para_lines)
        parts.append(f"<p>{_apply_inline_markdown(joined)}</p>")
        para_lines = []

    def close_ul() -> None:
        nonlocal in_ul
        if in_ul:
            parts.append("</ul>")
            in_ul = False

    for line in raw.splitlines():
        stripped = line.strip()
        if not stripped:
            flush_para()
            close_ul()
            continue

        if stripped.startswith("<") and (
            stripped.startswith("<div")
            or stripped.startswith("</div")
            or stripped.startswith("<span")
            or stripped.startswith("</span")
        ):
            flush_para()
            close_ul()
            parts.append(stripped)
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            flush_para()
            close_ul()
            level = len(heading.group(1))
            parts.append(f"<h{level}>{_apply_inline_markdown(heading.group(2))}</h{level}>")
            continue

        if re.match(r"^[-*]\s+", stripped):
            flush_para()
            if not in_ul:
                parts.append("<ul>")
                in_ul = True
            item = re.sub(r"^[-*]\s+", "", stripped)
            parts.append(f"<li>{_apply_inline_markdown(item)}</li>")
            continue

        flush_para()
        close_ul()
        para_lines.append(stripped)

    flush_para()
    close_ul()
    return "\n".join(parts) if parts else "<p></p>"


def _markdown_to_html_body(text: str) -> str:
    body = ""
    try:
        import markdown as md_lib

        body = md_lib.markdown(
            text or "",
            extensions=["extra", "tables", "fenced_code", "nl2br", "sane_lists"],
        )
    except ImportError:
        body = _basic_markdown_to_html_body(text)
    except Exception:
        try:
            import markdown as md_lib

            body = md_lib.markdown(
                text or "",
                extensions=["tables", "fenced_code", "nl2br"],
            )
        except Exception:
            body = _basic_markdown_to_html_body(text)
    return _sanitize_office_html(body)


def _markdown_inline_html(text: str) -> str:
    """Render one line/short block with emphasis (bold, italic) for slide bullets etc."""
    chunk = (text or "").strip()
    if not chunk:
        return ""
    try:
        import markdown as md_lib

        html = md_lib.markdown(
            chunk,
            extensions=["extra", "nl2br"],
            output_format="html5",
        )
        html = re.sub(r"^<p>", "", html)
        html = re.sub(r"</p>\s*$", "", html.strip())
        return _sanitize_office_html(html) or _apply_inline_markdown(chunk)
    except ImportError:
        return _apply_inline_markdown(chunk)
    except Exception:
        return _apply_inline_markdown(chunk)


def _parse_markdown_table_rows(markdown_text: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in (markdown_text or "").splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            if re.match(r"^\|\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|$", stripped):
                continue
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            rows.append(cells)
        elif "\t" in stripped:
            rows.append([c.strip() for c in stripped.split("\t")])
        elif "," in stripped and len(stripped.split(",")) > 1:
            rows.append([c.strip() for c in stripped.split(",")])
    return rows


def _rows_to_table_html(rows: list[list[str]]) -> str:
    if not rows:
        return f"<p>{_escape('No tabular data detected.')}</p>"
    parts = ['<table class="loma-grid-table">']
    for idx, row in enumerate(rows):
        tag = "th" if idx == 0 else "td"
        parts.append("<tr>")
        for cell in row:
            parts.append(f"<{tag}>{_escape(cell)}</{tag}>")
        parts.append("</tr>")
    parts.append("</table>")
    return "".join(parts)


def _wrap_paper_inner(body_html: str) -> str:
    return f'<div class="loma-paper-sheet"><div class="loma-paper-inner">{body_html}</div></div>'


def docx_to_markup_html(path: str) -> str:
    try:
        import mammoth
    except ImportError:
        from services.source_parser import parse_file

        parsed_source = parse_file(path)
        body = _markdown_to_html_body(parsed_source.markdown or parsed_source.legacy_upload_dict().get("content", ""))
        return _wrap_paper_inner(body)

    with open(path, "rb") as docx_file:
        result = mammoth.convert_to_html(docx_file)
    body = _sanitize_office_html(result.value or "<p></p>")
    return _wrap_paper_inner(body)


def _reflow_pdf_page_text(page_text: str) -> str:
    """Merge per-glyph line breaks (common in CJK PDFs) into horizontal paragraphs."""
    lines = [ln.strip() for ln in (page_text or "").splitlines() if ln.strip()]
    if not lines:
        return ""
    short = sum(1 for ln in lines if len(ln) <= 2)
    if len(lines) >= 8 and short / len(lines) >= 0.55:
        return "".join(lines)
    paragraphs: list[str] = []
    buf: list[str] = []
    for ln in lines:
        if len(ln) <= 2:
            buf.append(ln)
        else:
            if buf:
                paragraphs.append("".join(buf))
                buf = []
            paragraphs.append(ln)
    if buf:
        paragraphs.append("".join(buf))
    return "\n\n".join(paragraphs)


def pdf_to_plain_text(path: str) -> str:
    """Reflowed plain text for source/read-only panes."""
    try:
        import pypdf

        reader = pypdf.PdfReader(path)
        parts: list[str] = []
        for page_num, page in enumerate(reader.pages, start=1):
            page_text = _reflow_pdf_page_text((page.extract_text() or "").strip())
            if page_text:
                parts.append(f"--- Page {page_num} ---\n{page_text}")
        return "\n\n".join(parts) if parts else "(empty PDF)"
    except Exception as ex:
        return f"PDF read error: {ex}"


def pdf_to_markup_html(path: str) -> str:
    try:
        import pypdf

        reader = pypdf.PdfReader(path)
        pages: list[str] = []
        for page_num, page in enumerate(reader.pages, start=1):
            page_text = _reflow_pdf_page_text((page.extract_text() or "").strip())
            if not page_text:
                page_text = "(empty page)"
            paras = "".join(
                f"<p>{_escape(p)}</p>" for p in page_text.split("\n\n") if p.strip()
            ) or f"<p>{_escape(page_text)}</p>"
            pages.append(
                f'<div class="loma-paper-sheet">'
                f'<div class="loma-paper-inner">'
                f'<div class="loma-page-number">Page {page_num}</div>{paras}'
                f"</div></div>"
            )
        return "".join(pages) if pages else "<p>No pages found.</p>"
    except Exception as ex:
        return f"<p class='loma-markup-error'>PDF preview error: {_escape(str(ex))}</p>"


def _pptx_run_span(run) -> str:
    text = run.text or ""
    if not text:
        return ""
    styles: list[str] = []
    font = run.font
    if font.bold:
        styles.append("font-weight:700")
    if font.italic:
        styles.append("font-style:italic")
    if font.underline:
        styles.append("text-decoration:underline")
    try:
        if font.size and font.size.pt:
            styles.append(f"font-size:{round(font.size.pt, 1)}pt")
    except Exception:
        pass
    try:
        if font.color and font.color.rgb:
            rgb = font.color.rgb
            styles.append(f"color:#{rgb}")
    except Exception:
        pass
    try:
        if font.name:
            styles.append(f"font-family:{font.name},Segoe UI,system-ui,sans-serif")
    except Exception:
        pass
    inner = _escape(text).replace("\n", "<br/>")
    if styles:
        return f'<span style="{";".join(styles)}">{inner}</span>'
    return inner


def _pptx_paragraph_html(paragraph) -> str:
    parts = [_pptx_run_span(run) for run in paragraph.runs if (run.text or "").strip()]
    if parts:
        return "".join(parts)
    text = (paragraph.text or "").strip()
    return _escape(text) if text else ""


def _pptx_shape_box_style(shape, slide_w, slide_h) -> str:
    try:
        sw = int(slide_w) or 1
        sh = int(slide_h) or 1
        left = int(shape.left) / sw * 100
        top = int(shape.top) / sh * 100
        width = max(8.0, int(shape.width) / sw * 100)
        height = max(6.0, int(shape.height) / sh * 100)
        return (
            f"left:{left:.2f}%;top:{top:.2f}%;width:{width:.2f}%;height:{height:.2f}%;"
        )
    except Exception:
        return "left:6%;top:8%;width:88%;height:40%;"


def _pptx_text_frame_html(text_frame) -> str:
    parts: list[str] = []
    for para in text_frame.paragraphs:
        html_p = _pptx_paragraph_html(para)
        if not html_p:
            continue
        level = getattr(para, "level", 0) or 0
        cls = "loma-slide-body" if level == 0 else "loma-slide-bullet"
        margin = "" if level == 0 else f' style="margin-left:{level * 1.1}em"'
        parts.append(f'<p class="{cls}"{margin}>{html_p}</p>')
    return "".join(parts)


def _pptx_slide_background_style(slide) -> str:
    try:
        fill = slide.background.fill
        if fill.type is not None and hasattr(fill, "fore_color") and fill.fore_color.rgb:
            return f"background-color:#{fill.fore_color.rgb};"
    except Exception:
        pass
    return "background-color:#FFFAF5;"


def _pptx_font_size_cqh(pt: float, slide_h_emu: int) -> str:
    """Map PowerPoint point size to container-query height units on the slide frame."""
    slide_h_in = max(int(slide_h_emu) / 914400, 1)
    pct = (float(pt) / 72.0) / slide_h_in * 100.0
    return f"{pct:.3f}cqh"


def _pptx_paragraph_color(paragraph, *, is_title: bool = False) -> str:
    try:
        if paragraph.font.color and paragraph.font.color.rgb:
            return f"#{paragraph.font.color.rgb}"
    except Exception:
        pass
    return "#3E2723" if is_title else "#4E342E"


def _pptx_paragraph_size_pt(paragraph, *, is_title: bool = False) -> float:
    try:
        if paragraph.font.size and paragraph.font.size.pt:
            return float(paragraph.font.size.pt)
    except Exception:
        pass
    return 32.0 if is_title else 18.0


def _pptx_paragraph_html_scaled(paragraph, slide_h_emu: int, *, is_title: bool = False) -> str:
    parts = [_pptx_run_span(run) for run in paragraph.runs if (run.text or "").strip()]
    if parts:
        inner = "".join(parts)
    else:
        text = (paragraph.text or "").strip()
        inner = _escape(text) if text else ""
    if not inner:
        return ""
    pt = _pptx_paragraph_size_pt(paragraph, is_title=is_title)
    color = _pptx_paragraph_color(paragraph, is_title=is_title)
    size = _pptx_font_size_cqh(pt, slide_h_emu)
    styles = [f"color:{color}", f"font-size:{size}", "line-height:1.3", "margin:0 0 0.28em"]
    if is_title or (paragraph.font.bold if paragraph.font else False):
        styles.append("font-weight:700")
    return f'<p style="{";".join(styles)}">{inner}</p>'


def _pptx_text_frame_html_scaled(text_frame, slide_h_emu: int, *, is_title: bool = False) -> str:
    parts: list[str] = []
    for para in text_frame.paragraphs:
        html_p = _pptx_paragraph_html_scaled(para, slide_h_emu, is_title=is_title)
        if html_p:
            parts.append(html_p)
    return "".join(parts)


def _pptx_shape_fill_style(shape) -> str:
    try:
        if shape.fill.type is not None and shape.fill.fore_color.rgb:
            return f"background-color:#{shape.fill.fore_color.rgb};"
    except Exception:
        pass
    return ""


def _pptx_shape_sort_key(shape) -> tuple[int, int]:
    """Decorative shapes first, then text shapes top-to-bottom."""
    has_text = 0
    if shape.has_text_frame and (shape.text_frame.text or "").strip():
        has_text = 1
    try:
        top = int(shape.top)
    except Exception:
        top = 0
    return (has_text, top)


def pptx_to_markup_html(path: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return "<p class='loma-markup-error'>python-pptx is required for slide preview.</p>"

    prs = Presentation(path)
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    aspect = f"{slide_w} / {slide_h}"

    slides_html: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        blocks: list[str] = []
        title_shape = slide.shapes.title if slide.shapes.title else None
        for shape in sorted(slide.shapes, key=_pptx_shape_sort_key):
            box = _pptx_shape_box_style(shape, slide_w, slide_h)
            if shape.has_text_frame:
                text = (shape.text_frame.text or "").strip()
                if not text:
                    continue
                is_title = shape is title_shape
                inner_html = _pptx_text_frame_html_scaled(
                    shape.text_frame, slide_h, is_title=is_title
                )
                if not inner_html:
                    continue
                title_cls = " loma-slide-title" if is_title else " loma-slide-body"
                blocks.append(
                    f'<div class="loma-slide-shape{title_cls}" style="position:absolute;{box}">'
                    f"{inner_html}</div>"
                )
                continue
            fill = _pptx_shape_fill_style(shape)
            if fill:
                blocks.append(
                    f'<div class="loma-slide-decor" style="position:absolute;{box}{fill}"></div>'
                )

        bg = _pptx_slide_background_style(slide)
        inner = (
            "".join(blocks)
            if blocks
            else f"<p style='color:#3E2723;font-size:4cqh'>{_escape(f'Slide {idx}')}</p>"
        )
        slides_html.append(
            f'<div class="loma-slide-frame loma-slide-frame-pptx">'
            f'<div class="loma-slide-inner" style="{bg}">{inner}</div>'
            f"</div>"
        )
    deck = "".join(slides_html) if slides_html else "<p>No slides found.</p>"
    return (
        f'<div class="loma-slide-deck loma-slide-deck-pptx" '
        f'style="--slide-aspect: {aspect}">{deck}</div>'
    )


def xlsx_to_markup_html(path: str) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "<p class='loma-markup-error'>openpyxl is required for spreadsheet preview.</p>"

    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(["" if cell is None else str(cell) for cell in row])
    wb.close()
    table = _rows_to_table_html(rows)
    return f'<div class="loma-sheet-grid">{table}</div>'


def _draft_document_html(draft: str) -> str:
    body = _markdown_to_html_body(draft)
    return _wrap_paper_inner(body)


def _draft_presentation_html(draft: str) -> str:
    normalized = normalize_presentation_markdown(draft or "")
    blocks = split_presentation_slides(normalized)
    if not blocks:
        blocks = [normalized] if normalized else []
    slides: list[str] = []
    for idx, block in enumerate(blocks, start=1):
        lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
        title = lines[0].lstrip("#").strip() if lines else f"Slide {idx}"
        bullets = [ln.lstrip("*- ").strip() for ln in lines[1:] if ln.lstrip().startswith(("*", "-"))]
        if not bullets and len(lines) > 1:
            bullets = lines[1:]
        title_html = f"<h2 class='loma-slide-title'>{_markdown_inline_html(title)}</h2>"
        list_html = ""
        if bullets:
            items = "".join(f"<li>{_markdown_inline_html(b)}</li>" for b in bullets)
            list_html = f"<ul class='loma-slide-bullets'>{items}</ul>"
        elif len(lines) > 1:
            list_html = _markdown_to_html_body("\n\n".join(lines[1:]))
        slides.append(
            f'<div class="loma-slide-frame loma-slide-frame-draft">'
            f'<div class="loma-slide-inner">{title_html}{list_html}</div>'
            f"</div>"
        )
    deck = "".join(slides) if slides else f"<p>{_escape('No slide content.')}</p>"
    return f'<div class="loma-slide-deck loma-slide-deck-draft">{deck}</div>'


def _draft_spreadsheet_html(draft: str) -> str:
    rows = _parse_markdown_table_rows(draft)
    if not rows and (draft or "").strip():
        rows = [[(draft or "").strip()[:32000]]]
    table = _rows_to_table_html(rows)
    return f'<div class="loma-sheet-grid">{table}</div>'


def draft_to_markup_html(draft: str, output_type: str) -> str:
    text = (draft or "").strip()
    out = (output_type or "document").lower()
    if out == "presentation":
        return _draft_presentation_html(text)
    if out == "spreadsheet":
        return _draft_spreadsheet_html(text)
    if out in ("chat", "software", "sound", "video", "image"):
        body = _markdown_to_html_body(text) if text else "<p></p>"
        return f'<div class="loma-markup-plain">{body}</div>'
    return _draft_document_html(text)


def preview_html_from_draft(draft: str, output_type: str) -> str:
    """GitHub-style HTML preview from markdown source only (Preview panel Viewer mode)."""
    text = (draft or "").strip()
    if not text:
        out = (output_type or "document").lower()
        if out == "presentation":
            return '<div class="loma-slide-deck"><p></p></div>'
        if out == "spreadsheet":
            return '<div class="loma-sheet-grid"><p></p></div>'
        return '<div class="loma-paper-sheet"><div class="loma-paper-inner"><p></p></div></div>'
    text = re.sub(
        r"\[IMAGE_PROMPT:\s*(.+?)\]",
        r'\n\n<div class="loma-image-placeholder">Image: \1</div>\n\n',
        text,
        flags=re.IGNORECASE,
    )
    return draft_to_markup_html(text, output_type)


def _path_matches_output(path: str, output_type: str) -> bool:
    ext = os.path.splitext(path)[1].lower()
    expected = _EXTENSION_BY_OUTPUT.get(output_type.lower())
    if expected and ext == expected:
        return True
    if output_type.lower() == "pdf" and ext == ".pdf":
        return True
    return False


def resolve_markup_html(*, path: str | None, draft: str, output_type: str) -> str:
    """Prefer compiled/uploaded file when extension matches; else draft sandbox."""
    if path and os.path.isfile(path):
        ext = os.path.splitext(path)[1].lower()
        out = (output_type or "document").lower()
        if ext == ".docx" and out in ("document", "docx"):
            return docx_to_markup_html(path)
        if ext == ".pdf" and out in ("pdf", "document"):
            return pdf_to_markup_html(path)
        if ext == ".pptx" and out == "presentation":
            return pptx_to_markup_html(path)
        if ext == ".xlsx" and out == "spreadsheet":
            return xlsx_to_markup_html(path)
        if _path_matches_output(path, out):
            if ext == ".docx":
                return docx_to_markup_html(path)
            if ext == ".pdf":
                return pdf_to_markup_html(path)
            if ext == ".pptx":
                return pptx_to_markup_html(path)
            if ext == ".xlsx":
                return xlsx_to_markup_html(path)
    return draft_to_markup_html(draft, output_type)


def resolve_document_viewer_html(path: str, markdown_source: str) -> str:
    """Viewer HTML from extracted markdown source (single source of truth)."""
    _ = path
    return preview_html_from_draft(markdown_source, "document")
